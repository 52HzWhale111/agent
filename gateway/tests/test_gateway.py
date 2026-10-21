import asyncio
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import time
import zipfile

import httpx
import pytest
from fastapi.testclient import TestClient
from gateway.main import create_app


def finished(client, task_id):
    for _ in range(100):
        result = client.get('/api/tasks/' + task_id).json()
        if result['status'] in ('succeeded', 'failed'):
            return result
        time.sleep(.03)
    raise AssertionError('Task never finished')


def archive(name='Main.java', content='class Main {}'):
    data = BytesIO()
    with zipfile.ZipFile(data, 'w') as output:
        output.writestr(name, content)
    return data.getvalue()


def test_pr2_routes_and_durable_tasks(tmp_path, monkeypatch):
    monkeypatch.delenv('GATEWAY_API_KEY', raising=False)
    calls = []
    def agent(request):
        calls.append((request.url.path, json.loads(request.content), request.headers['x-request-id']))
        return httpx.Response(200, json={'answer': 'ok', 'sources': ['chapter 5']})
    app = create_app(tmp_path, httpx.MockTransport(agent))
    with TestClient(app) as client:
        assert client.post('/api/qa', json={'question': 'cohesion'}, headers={'X-Request-ID': 'trace-1'}).status_code == 200
        assert client.post('/api/requirements/analyze', json={'description': 'library'}).status_code == 200
        review = client.post('/api/reviews', json={'content': 'SRS text', 'doc_type': 'srs'})
        assert review.status_code == 202
        result = finished(client, review.json()['task_id'])
        assert result['status'] == 'succeeded' and result['progress'] == 100
        assert result['result']['answer'] == 'ok'
        code = client.post('/api/code-tests', json={'code': 'class Main {}'})
        assert finished(client, code.json()['task_id'])['status'] == 'succeeded'
        assert client.post('/api/consistency', json={'code': 'class Main {}', 'diagram': {}}).status_code == 200
        assert [entry[0] for entry in calls] == ['/qa', '/analyze_requirement', '/review_document', '/run_tests', '/check_consistency']
        assert calls[0][2] == 'trace-1'
        assert calls[2][1] == {'content': 'SRS text', 'doc_type': 'srs'}
    with TestClient(create_app(tmp_path, httpx.MockTransport(agent))) as client:
        assert client.get('/api/tasks/' + result['task_id']).json()['status'] == 'succeeded'
    assert 'trace-1' in (tmp_path / 'logs/gateway.log').read_text(encoding='utf-8')


def test_timeouts_failure_isolation_and_health(tmp_path, monkeypatch):
    monkeypatch.setenv('QA_TIMEOUT', '.02')
    monkeypatch.setenv('DOC_REVIEW_TIMEOUT', '.02')
    monkeypatch.setenv('CODE_RUN_TIMEOUT', '.02')
    async def agent(request):
        if request.url.host == 'localhost' and request.url.port == 8200:
            raise httpx.ConnectError('offline')
        if request.url.path == '/health':
            return httpx.Response(200, json={'status': 'ok'})
        await asyncio.sleep(.1)
        return httpx.Response(200, json={})
    with TestClient(create_app(tmp_path, httpx.MockTransport(agent))) as client:
        response = client.post('/api/qa', json={'question': 'q'})
        assert response.status_code == 504
        assert response.json()['error']['code'] == 'AGENT_TIMEOUT'
        review = client.post('/api/reviews', json={'content': 'text', 'doc_type': 'srs'})
        assert finished(client, review.json()['task_id'])['error']['code'] == 'AGENT_TIMEOUT'
        code = client.post('/api/code-tests', json={'code': 'class Main {}'})
        assert finished(client, code.json()['task_id'])['error']['code'] == 'AGENT_UNAVAILABLE'
        ready = client.get('/health/ready')
        assert ready.status_code == 503
        assert ready.json()['services'] == {'knowledge': 'ok', 'sandbox': 'unavailable'}
        assert client.get('/health/live').status_code == 200


def test_upload_limits_random_names_and_archive_handoff(tmp_path, monkeypatch):
    monkeypatch.setenv('MAX_UPLOAD_MB', '1')
    def agent(request):
        payload = json.loads(request.content)
        assert Path(payload['archive_path']).is_file()
        return httpx.Response(200, json={'passed': 1})
    with TestClient(create_app(tmp_path, httpx.MockTransport(agent))) as client:
        for filename, data in [('a.exe', b'bad'), ('a.pdf', b'not pdf'), ('a.docx', archive())]:
            assert client.post('/api/uploads', files={'file': (filename, data)}).status_code == 415
        large = client.post('/api/uploads', files={'file': ('a.pdf', b'%PDF-' + b'x' * (1024 * 1024))})
        assert large.status_code == 413
        assert list((tmp_path / 'uploads').iterdir()) == []
        first = client.post('/api/uploads', files={'file': ('../../same.zip', archive())}).json()
        second = client.post('/api/uploads', files={'file': ('same.zip', archive())}).json()
        assert first['file_id'] != second['file_id']
        assert len(list((tmp_path / 'uploads').iterdir())) == 2
        task = client.post('/api/code-tests', json={'file_id': first['file_id']})
        assert finished(client, task.json()['task_id'])['result'] == {'passed': 1}
        assert client.post('/api/reviews', json={'file_id': first['file_id'], 'doc_type': 'srs'}).status_code == 422
        assert client.post('/api/code-tests', json={'file_id': '../../missing'}).status_code == 404


def test_docx_text_and_table_extraction(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph('Requirements')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Acceptance criteria'
    data = BytesIO()
    doc.save(data)
    def agent(request):
        body = json.loads(request.content)
        assert 'Requirements' in body['content']
        assert 'Acceptance criteria' in body['content']
        return httpx.Response(200, json={'issues': []})
    with TestClient(create_app(tmp_path, httpx.MockTransport(agent))) as client:
        uploaded = client.post('/api/uploads', files={'file': ('requirements.docx', data.getvalue())}).json()
        task = client.post('/api/reviews', json={'file_id': uploaded['file_id'], 'doc_type': 'srs'})
        assert task.status_code == 202
        assert finished(client, task.json()['task_id'])['status'] == 'succeeded'


def test_validation_auth_and_queue_limit(tmp_path, monkeypatch):
    monkeypatch.setenv('GATEWAY_API_KEY', 'local-test-secret')
    monkeypatch.setenv('MAX_PENDING_TASKS', '1')
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        assert client.post('/api/qa', json={'question': 'q'}).status_code == 401
        assert client.get('/health/live').status_code == 200
        client.headers['Authorization'] = 'Bearer local-test-secret'
        bad = client.post('/api/qa', json={'question': ''})
        assert bad.status_code == 422 and 'question' not in str(bad.json())
        assert client.post('/api/reviews', json={'doc_type': 'srs'}).status_code == 422
        assert client.post('/api/reviews', json={'content': 'text', 'file_id': 'x', 'doc_type': 'srs'}).status_code == 422
        assert client.get('/api/tasks/missing').status_code == 404
        first = client.post('/api/code-tests', json={'code': 'class Main {}'})
        assert first.status_code == 202
        assert client.post('/api/code-tests', json={'code': 'class Main {}'}).status_code == 503


def test_queued_work_survives_restart_and_interrupted_work_fails(tmp_path):
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        queued = client.post('/api/code-tests', json={'code': 'class Main {}'}).json()['task_id']
        running = client.post('/api/code-tests', json={'code': 'class Main {}'}).json()['task_id']
    with sqlite3.connect(tmp_path / 'gateway.sqlite3') as conn:
        conn.execute("UPDATE tasks SET status='running' WHERE id=?", (running,))
    transport = httpx.MockTransport(lambda request: httpx.Response(200, json={'passed': 1}))
    with TestClient(create_app(tmp_path, transport)) as client:
        assert finished(client, queued)['status'] == 'succeeded'
        assert finished(client, running)['error']['code'] == 'WORKER_RESTARTED'


@pytest.mark.parametrize('status,body,code', [(503, {}, 'AGENT_UNAVAILABLE'), (400, {}, 'AGENT_REJECTED'), (302, {}, 'AGENT_REJECTED'), (200, [], 'INVALID_AGENT_RESPONSE')])
def test_agent_errors(tmp_path, status, body, code):
    transport = httpx.MockTransport(lambda request: httpx.Response(status, json=body))
    with TestClient(create_app(tmp_path, transport)) as client:
        response = client.post('/api/qa', json={'question': 'q'})
        assert response.json()['error']['code'] == code


def test_code_timeout_and_exact_upload_limit(tmp_path, monkeypatch):
    monkeypatch.setenv('CODE_RUN_TIMEOUT', '.02')
    monkeypatch.setenv('MAX_UPLOAD_MB', '20')
    async def slow(request):
        await asyncio.sleep(.1)
        return httpx.Response(200, json={})
    with TestClient(create_app(tmp_path, httpx.MockTransport(slow))) as client:
        body = b'%PDF-' + b'x' * (20 * 1024 * 1024 - 5)
        exact = client.post('/api/uploads', files={'file': ('boundary.pdf', body)})
        assert exact.status_code == 201
        assert exact.json()['size'] == 20 * 1024 * 1024
        over = client.post('/api/uploads', files={'file': ('over.pdf', body + b'x')})
        assert over.status_code == 413
        task = client.post('/api/code-tests', json={'code': 'class Main {}'}).json()
        assert finished(client, task['task_id'])['error']['code'] == 'AGENT_TIMEOUT'


def test_bad_document_task_does_not_stop_worker(tmp_path):
    transport = httpx.MockTransport(lambda request: httpx.Response(200, json={'passed': 1}))
    with TestClient(create_app(tmp_path, transport)) as client:
        file = client.post('/api/uploads', files={'file': ('broken.pdf', b'%PDF-broken')}).json()
        task = client.post('/api/reviews', json={'file_id': file['file_id'], 'doc_type': 'srs'}).json()
        assert finished(client, task['task_id'])['error']['code'] == 'DOCUMENT_PARSE_FAILED'
        code = client.post('/api/code-tests', json={'code': 'class Main {}'}).json()
        assert finished(client, code['task_id'])['status'] == 'succeeded'


def test_missing_upload_is_404_before_enqueue(tmp_path):
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        uploaded = client.post('/api/uploads', files={'file': ('project.zip', archive())}).json()
        (tmp_path / 'uploads' / (uploaded['file_id'] + '.zip')).unlink()
        response = client.post('/api/code-tests', json={'file_id': uploaded['file_id']})
        assert response.status_code == 404 and response.json()['error']['code'] == 'FILE_NOT_FOUND'


def test_non_ascii_credentials_return_401(tmp_path, monkeypatch):
    monkeypatch.setenv('GATEWAY_API_KEY', 'secret')
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        response = client.get('/api/tasks/missing', headers={b'Authorization': b'Bearer \xc3\xa9'})
        assert response.status_code == 401


def test_swagger_authorize_with_protected_api(tmp_path, monkeypatch):
    monkeypatch.setenv('GATEWAY_API_KEY', 'test-secret')
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        assert client.get('/docs').status_code == 200
        schema = client.get('/openapi.json').json()
        assert schema['components']['securitySchemes']['HTTPBearer']['scheme'] == 'bearer'
        assert client.get('/api/tasks/missing').status_code == 401
        response = client.get('/api/tasks/missing', headers={'Authorization': 'Bearer test-secret'})
        assert response.status_code == 404
