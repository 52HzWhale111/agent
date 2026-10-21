"""Single-host Agent gateway; SQLite is the durable task queue."""
import asyncio
from contextlib import asynccontextmanager, contextmanager
from datetime import datetime, timezone
import json
import logging
from logging.handlers import TimedRotatingFileHandler
import os
from pathlib import Path
import re
import secrets
import sqlite3
import time
from typing import Literal
from urllib.parse import quote
import uuid
import zipfile

import httpx
from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse, Response
from fastapi.security import HTTPBearer
from pydantic import BaseModel, Field


class QAQuery(BaseModel):
    question: str = Field(min_length=1, max_length=20000)
    user_id: str | None = None


class RequirementQuery(BaseModel):
    description: str = Field(min_length=1, max_length=30000)


class ReviewQuery(BaseModel):
    content: str | None = Field(default=None, max_length=500000)
    file_id: str | None = None
    document_id: str | None = None
    doc_type: Literal["srs", "hld", "test_report"]
    scope: Literal["auto", "full", "excerpt"] | None = None


class CodeQuery(BaseModel):
    code: str | None = Field(default=None, max_length=200000)
    file_id: str | None = None
    language: Literal["java"] = "java"
    tests: str = Field(default="", max_length=200000)


class ConsistencyQuery(BaseModel):
    code: str = Field(min_length=1, max_length=200000)
    diagram: dict


def create_app(data_dir=None, transport=None, start_worker=True):
    root = Path(data_dir or os.getenv("GATEWAY_DATA_DIR", "./data"))
    uploads = root / "uploads"
    uploads.mkdir(parents=True, exist_ok=True)
    logs = root / "logs"
    logs.mkdir(exist_ok=True)
    logger = logging.getLogger(f"gateway.{root.resolve()}")
    logger.setLevel(logging.INFO)
    if not logger.handlers:
        handler = TimedRotatingFileHandler(logs / "gateway.log", when="midnight", backupCount=14, utc=True, encoding="utf-8")
        logger.addHandler(handler)
        logger.addHandler(logging.StreamHandler())
    db_path = root / "gateway.sqlite3"
    services = {
        "knowledge": os.getenv("KNOWLEDGE_AGENT_URL", "http://localhost:8100").rstrip("/"),
        "sandbox": os.getenv("SANDBOX_AGENT_URL", "http://localhost:8200").rstrip("/"),
    }
    backend_url = os.getenv("BACKEND_URL", "").rstrip("/")
    sandbox_protocol = os.getenv("SANDBOX_PROTOCOL", "text")
    if backend_url:
        services["backend"] = backend_url
    routes = {
        "qa": ("knowledge", "/qa", float(os.getenv("QA_TIMEOUT", "30"))),
        "requirement": ("knowledge", "/analyze_requirement", float(os.getenv("QA_TIMEOUT", "30"))),
        "review": ("knowledge", "/review_document", float(os.getenv("DOC_REVIEW_TIMEOUT", "60"))),
        "code": ("sandbox", "/sandbox/run-tests" if sandbox_protocol == "text" else "/run_tests", float(os.getenv("CODE_RUN_TIMEOUT", "120"))),
        "consistency": ("sandbox", "/check_consistency", float(os.getenv("QA_TIMEOUT", "30"))),
    }
    api_key = os.getenv("GATEWAY_API_KEY", "")
    max_upload = min(int(os.getenv("MAX_UPLOAD_MB", "20")), 20) * 1024 * 1024
    max_queue = int(os.getenv("MAX_PENDING_TASKS", "100"))

    @contextmanager
    def db():
        conn = sqlite3.connect(db_path, timeout=10)
        conn.row_factory = sqlite3.Row
        try:
            with conn:
                yield conn
        finally:
            conn.close()

    with db() as conn:
        conn.executescript("""
            PRAGMA journal_mode=WAL;
            CREATE TABLE IF NOT EXISTS files (
                id TEXT PRIMARY KEY, suffix TEXT NOT NULL, size INTEGER NOT NULL);
            CREATE TABLE IF NOT EXISTS tasks (
                id TEXT PRIMARY KEY, kind TEXT NOT NULL, payload TEXT NOT NULL,
                status TEXT NOT NULL, progress INTEGER NOT NULL DEFAULT 0,
                result TEXT, error TEXT, request_id TEXT NOT NULL, created_at TEXT NOT NULL);
        """)
        for table in ("files", "tasks"):
            if "owner_id" not in {row[1] for row in conn.execute(f"PRAGMA table_info({table})")}:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN owner_id TEXT")

    def log(**values):
        logger.info(json.dumps({"time": datetime.now(timezone.utc).isoformat(), **values}, ensure_ascii=False))

    def error(status, code, message):
        return HTTPException(status, {"code": code, "message": message})

    async def backend_request(method, path, request, **kwargs):
        headers = {"Authorization": request.headers.get("Authorization", ""),
                   "X-Request-ID": request.state.request_id, **kwargs.pop("headers", {})}
        try:
            response = await app.state.client.request(method, backend_url + path, headers=headers, timeout=15, **kwargs)
        except httpx.TimeoutException:
            raise error(504, "BACKEND_TIMEOUT", "业务服务响应较慢，请重试")
        except httpx.RequestError:
            raise error(503, "BACKEND_UNAVAILABLE", "业务服务暂不可用")
        log(event="backend_call", request_id=request.state.request_id,
            backend_request_id=response.headers.get("X-Request-ID"), path=path, status=response.status_code)
        return response

    def backend_data(response):
        try:
            body = response.json()
            if response.is_error:
                raise error(response.status_code, body.get("code", "BACKEND_REJECTED"), body.get("message", "业务请求失败"))
            if not response.is_success or not isinstance(body.get("data"), dict):
                raise ValueError()
            return body["data"]
        except (ValueError, AttributeError):
            raise error(502, "INVALID_BACKEND_RESPONSE", "业务服务返回格式异常")

    async def call_agent(kind, payload, request_id):
        service, path, timeout = routes[kind]
        started = time.monotonic()
        try:
            if kind == "consistency" and sandbox_protocol == "text":
                raise error(503, "FEATURE_UNAVAILABLE", "工程服务尚未提供图代码一致性接口")
            if kind == "code" and sandbox_protocol == "text":
                if payload.get("archive_path") or not payload.get("code", "").strip() or not payload.get("tests", "").strip():
                    raise error(422, "CODE_AND_TESTS_REQUIRED", "当前工程接口需要单份Java代码及JUnit测试源码")
                if "---TEST---" in payload["code"] or "---TEST---" in payload["tests"]:
                    raise error(422, "INVALID_INPUT", "源码不能包含工程接口保留分隔符 ---TEST---")
                arguments = {"content": (payload["code"] + "\n---TEST---\n" + payload["tests"]).encode("utf-8"),
                             "headers": {"X-Request-ID": request_id, "Content-Type": "text/plain; charset=utf-8"}}
            else:
                arguments = {"json": payload, "headers": {"X-Request-ID": request_id}}
            # wait_for limits the whole request, including a slowly streaming response.
            response = await asyncio.wait_for(
                app.state.client.post(services[service] + path, timeout=timeout, **arguments), timeout)
            if response.status_code >= 500:
                raise error(503, "AGENT_UNAVAILABLE", "该功能暂不可用")
            if not response.is_success:
                raise error(502, "AGENT_REJECTED", "Agent 未能处理该请求，请检查输入后重试")
            try:
                result = response.json()
                if not isinstance(result, dict):
                    raise ValueError("Expected object")
            except ValueError:
                raise error(502, "INVALID_AGENT_RESPONSE", "Agent 返回格式异常")
            if kind == "code" and sandbox_protocol == "text" and result.get("status") == "ERROR":
                raise error(503, "SANDBOX_ERROR", "工程服务执行异常，请联系维护人员")
            return result
        except (httpx.TimeoutException, asyncio.TimeoutError):
            raise error(504, "AGENT_TIMEOUT", "这次有点慢，请重试")
        except httpx.RequestError:
            raise error(503, "AGENT_UNAVAILABLE", "该功能暂不可用")
        finally:
            log(event="agent_call", request_id=request_id, step=kind,
                duration_ms=round((time.monotonic() - started) * 1000))

    async def worker():
        # One worker process only. Interrupted work is terminal: do not run student code twice.
        with db() as conn:
            conn.execute("UPDATE tasks SET status='failed', error=?, payload='{}' WHERE status='running'",
                         (json.dumps({"code": "WORKER_RESTARTED", "message": "服务已重启，请重新提交任务"}, ensure_ascii=False),))
        while True:
            with db() as conn:
                conn.execute("BEGIN IMMEDIATE")
                row = conn.execute("SELECT * FROM tasks WHERE status='queued' ORDER BY created_at LIMIT 1").fetchone()
                if row:
                    conn.execute("UPDATE tasks SET status='running', progress=10 WHERE id=?", (row["id"],))
            if not row:
                await asyncio.sleep(0.2)
                continue
            try:
                payload = json.loads(row["payload"])
                async def execute():
                    if row["kind"] == "review" and "file_id" in payload:
                        path = file_path(payload.pop("file_id"), row["owner_id"])
                        payload["content"] = await asyncio.to_thread(extract_text, path)
                    return await call_agent(row["kind"], payload, row["request_id"])
                try:
                    result = await asyncio.wait_for(execute(), routes[row["kind"]][2])
                except asyncio.TimeoutError:
                    raise error(504, "AGENT_TIMEOUT", "这次有点慢，请重试")
                with db() as conn:
                    conn.execute("UPDATE tasks SET status='succeeded', progress=100, result=?, payload='{}' WHERE id=?",
                                 (json.dumps(result, ensure_ascii=False), row["id"]))
                log(event="task_finished", request_id=row["request_id"], task_id=row["id"], status="succeeded")
            except HTTPException as exc:
                with db() as conn:
                    conn.execute("UPDATE tasks SET status='failed', error=?, payload='{}' WHERE id=?",
                                 (json.dumps(exc.detail, ensure_ascii=False), row["id"]))
                log(event="task_finished", request_id=row["request_id"], task_id=row["id"], status="failed", error=exc.detail)
            except Exception:
                with db() as conn:
                    conn.execute("UPDATE tasks SET status='failed', error=?, payload='{}' WHERE id=?",
                                 (json.dumps({"code": "TASK_FAILED", "message": "任务处理失败，请重试"}, ensure_ascii=False), row["id"]))
                log(event="task_finished", request_id=row["request_id"], task_id=row["id"], status="failed", error="TASK_FAILED")

    @asynccontextmanager
    async def lifespan(app):
        app.state.client = httpx.AsyncClient(transport=transport, follow_redirects=False)
        task = asyncio.create_task(worker()) if start_worker else None
        yield
        if task:
            task.cancel()
            try:
                await task
            except asyncio.CancelledError:
                pass
        await app.state.client.aclose()

    app = FastAPI(title="知构 Agent 接入层", version="1.0.0", lifespan=lifespan,
                  dependencies=[Depends(HTTPBearer(auto_error=False))])
    app.state.db_path = db_path

    @app.middleware("http")
    async def request_context(request: Request, call_next):
        incoming = request.headers.get("X-Request-ID", "")
        request.state.request_id = incoming if re.fullmatch(r"[a-zA-Z0-9_-]{1,64}", incoming) else uuid.uuid4().hex
        request.state.user_id = None
        started = time.monotonic()
        public_paths = {"/health/live", "/docs", "/docs/oauth2-redirect", "/openapi.json", "/redoc"}
        is_backend_proxy = request.url.path.startswith(("/api/v1/", "/business/"))
        if backend_url and request.url.path.startswith("/api/") and not is_backend_proxy:
            try:
                if not request.headers.get("Authorization", "").startswith("Bearer "):
                    raise error(401, "AUTH_REQUIRED", "请先登录")
                user = backend_data(await backend_request("GET", "/api/v1/auth/me", request))
                if not isinstance(user.get("id"), str) or not user["id"]:
                    raise error(502, "INVALID_BACKEND_RESPONSE", "业务服务返回格式异常")
                request.state.user_id = user["id"]
            except HTTPException as exc:
                response = JSONResponse({"error": exc.detail, "request_id": request.state.request_id}, status_code=exc.status_code)
            else:
                response = await call_next(request)
        elif request.url.path not in public_paths and not is_backend_proxy and api_key and not secrets.compare_digest(request.headers.get("Authorization", "").encode("utf-8"), f"Bearer {api_key}".encode("utf-8")):
            response = JSONResponse({"error": {"code": "UNAUTHORIZED", "message": "请提供有效的访问凭证"},
                                     "request_id": request.state.request_id}, status_code=401)
        else:
            response = await call_next(request)
        response.headers["X-Request-ID"] = request.state.request_id
        log(event="request", request_id=request.state.request_id, method=request.method, path=request.url.path,
            status=response.status_code, duration_ms=round((time.monotonic() - started) * 1000))
        return response

    @app.exception_handler(HTTPException)
    async def handle_error(request, exc):
        detail = exc.detail if isinstance(exc.detail, dict) else {"code": "REQUEST_ERROR", "message": str(exc.detail)}
        return JSONResponse({"error": detail, "request_id": request.state.request_id}, status_code=exc.status_code)

    @app.exception_handler(RequestValidationError)
    async def handle_validation(request, exc):
        return JSONResponse({"error": {"code": "INVALID_INPUT", "message": "请求参数不正确"},
                             "request_id": request.state.request_id}, status_code=422)

    def enqueue(kind, payload, request_id, owner_id=None):
        task_id = uuid.uuid4().hex
        with db() as conn:
            conn.execute("BEGIN IMMEDIATE")
            count = conn.execute("SELECT count(*) FROM tasks WHERE status IN ('queued','running')").fetchone()[0]
            if count >= max_queue:
                raise error(503, "QUEUE_FULL", "任务较多，请稍后重试")
            conn.execute("INSERT INTO tasks(id,kind,payload,status,request_id,created_at,owner_id) VALUES (?,?,?,'queued',?,?,?)",
                         (task_id, kind, json.dumps(payload, ensure_ascii=False), request_id, datetime.now(timezone.utc).isoformat(), owner_id))
        return {"task_id": task_id, "status": "queued", "progress": 0, "poll_after_seconds": 2}

    def file_path(file_id, owner_id=None):
        with db() as conn:
            row = conn.execute("SELECT * FROM files WHERE id=?", (file_id,)).fetchone()
        if not row or (backend_url and row["owner_id"] != owner_id) or not (uploads / (row["id"] + row["suffix"])).is_file():
            raise error(404, "FILE_NOT_FOUND", "文件不存在")
        return uploads / (row["id"] + row["suffix"])

    @app.get("/health/live")
    async def live():
        return {"status": "ok", "service": "gateway"}

    @app.get("/health/ready")
    async def ready():
        async def check(name, url):
            try:
                if name == "sandbox" and sandbox_protocol == "text":
                    response = await asyncio.wait_for(app.state.client.options(url + "/sandbox/run-tests", timeout=2), 2)
                    return name, "http_route_only" if response.status_code == 204 else "unavailable"
                response = await asyncio.wait_for(app.state.client.get(url + "/health", timeout=2), 2)
                return name, "ok" if response.is_success else "unavailable"
            except (httpx.RequestError, asyncio.TimeoutError):
                return name, "unavailable"
        states = dict(await asyncio.gather(*(check(name, url) for name, url in services.items())))
        healthy = all(value in {"ok", "http_route_only"} for value in states.values())
        return JSONResponse({"status": "ok" if healthy else "degraded", "services": states}, status_code=200 if healthy else 503)

    @app.post("/api/qa")
    async def qa(query: QAQuery, request: Request):
        payload = query.model_dump(exclude_none=True)
        if backend_url:
            payload["user_id"] = request.state.user_id
        return await call_agent("qa", payload, request.state.request_id)

    @app.post("/api/requirements/analyze")
    async def requirement(query: RequirementQuery, request: Request):
        return await call_agent("requirement", query.model_dump(), request.state.request_id)

    @app.post("/api/consistency")
    async def consistency(query: ConsistencyQuery, request: Request):
        return await call_agent("consistency", query.model_dump(), request.state.request_id)

    @app.api_route("/api/v1/{path:path}", methods=["GET", "POST", "PUT", "OPTIONS"], include_in_schema=False)
    @app.api_route("/business/{path:path}", methods=["GET", "POST", "PUT", "OPTIONS"], include_in_schema=False)
    async def business_proxy(path: str, request: Request):
        if not backend_url:
            raise error(503, "BACKEND_UNAVAILABLE", "业务服务尚未配置")
        body = bytearray()
        async for chunk in request.stream():
            body.extend(chunk)
            if len(body) > 21 * 1024 * 1024:
                raise error(413, "FILE_TOO_LARGE", "请求不能超过21MiB")
        upstream_path = "/" + path if request.url.path.startswith("/business/") else "/api/v1/" + path
        response = await backend_request(request.method, upstream_path, request, content=bytes(body),
                                         params=request.query_params,
                                         headers={"Content-Type": request.headers.get("Content-Type", "application/json")})
        headers = {key: value for key, value in response.headers.items() if key.lower() in {"content-type", "content-disposition"}}
        if response.headers.get("X-Request-ID"):
            headers["X-Backend-Request-ID"] = response.headers["X-Request-ID"]
        return Response(response.content, status_code=response.status_code, headers=headers)

    @app.post("/api/uploads", status_code=201)
    async def upload(request: Request, file: UploadFile = File(...)):
        if backend_url:
            try:
                metadata = backend_data(await backend_request("POST", "/api/v1/documents", request,
                    files={"file": (file.filename, file.file, file.content_type or "application/octet-stream")}))
            finally:
                await file.close()
            return {"file_id": metadata["id"], "document_id": metadata["id"], "size": metadata["size_bytes"],
                    "extension": Path(metadata["original_name"]).suffix.lower(), "document": metadata}
        suffix = Path(file.filename or "").suffix.lower()
        if suffix not in {".doc", ".docx", ".pdf", ".zip"}:
            raise error(415, "FILE_TYPE_NOT_ALLOWED", "仅支持 doc/docx/pdf/zip")
        file_id = uuid.uuid4().hex
        target = uploads / (file_id + suffix)
        size = 0
        try:
            with target.open("xb") as output:
                while chunk := await file.read(65536):
                    size += len(chunk)
                    if size > max_upload:
                        raise error(413, "FILE_TOO_LARGE", "文件不能超过20MB")
                    output.write(chunk)
            with target.open("rb") as content:
                magic = content.read(8)
            valid = (suffix == ".pdf" and magic.startswith(b"%PDF-")) or (suffix == ".doc" and magic == bytes.fromhex("D0CF11E0A1B11AE1"))
            if suffix in {".docx", ".zip"} and zipfile.is_zipfile(target):
                with zipfile.ZipFile(target) as archive:
                    valid = suffix == ".zip" or "word/document.xml" in archive.namelist()
            if not valid:
                raise error(415, "INVALID_FILE_CONTENT", "文件内容与扩展名不匹配")
            with db() as conn:
                conn.execute("INSERT INTO files(id,suffix,size,owner_id) VALUES (?,?,?,?)", (file_id, suffix, size, request.state.user_id))
        except BaseException:
            target.unlink(missing_ok=True)
            raise
        finally:
            await file.close()
        return {"file_id": file_id, "size": size, "extension": suffix}

    def extract_text(path):
        try:
            if path.suffix == ".pdf":
                from pypdf import PdfReader
                result = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
            elif path.suffix == ".docx":
                from docx import Document
                with zipfile.ZipFile(path) as archive:
                    if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                        raise ValueError("Expanded document too large")
                document = Document(path)
                result = "\n".join([p.text for p in document.paragraphs] +
                                   [" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows])
            else:
                raise error(422, "TEXT_REQUIRED", "doc 请先转为 docx/pdf，zip 请用于代码测试；也可提交提取后的 content")
            if not result.strip() or len(result) > 500000:
                raise ValueError("Empty or oversized text")
            return result
        except HTTPException:
            raise
        except Exception:
            raise error(422, "DOCUMENT_PARSE_FAILED", "无法提取文档正文，请提交可复制文字的 docx/pdf 或 content")

    @app.post("/api/reviews", status_code=202)
    async def review(query: ReviewQuery, request: Request):
        if sum(bool(value) for value in (query.content, query.file_id, query.document_id)) != 1:
            raise error(422, "INVALID_INPUT", "content、file_id、document_id 必须且只能提供一个")
        payload = query.model_dump(exclude_none=True, exclude={"document_id"})
        reference = query.document_id or query.file_id
        if backend_url and reference:
            route = "/api/v1/documents/" + quote(reference, safe="")
            metadata = backend_data(await backend_request("GET", route, request))
            suffix = Path(metadata["original_name"]).suffix.lower()
            if suffix not in {".docx", ".pdf"}:
                raise error(422, "TEXT_REQUIRED", "评审请提交可提取正文的docx/pdf")
            if metadata.get("size_bytes", max_upload + 1) > max_upload:
                raise error(413, "FILE_TOO_LARGE", "文件不能超过20MiB")
            downloaded = await backend_request("GET", route + "/content", request)
            if not downloaded.is_success:
                backend_data(downloaded)
            if len(downloaded.content) > max_upload:
                raise error(413, "FILE_TOO_LARGE", "文件不能超过20MiB")
            cache_id = uuid.uuid4().hex
            cached = uploads / (cache_id + suffix)
            try:
                cached.write_bytes(downloaded.content)
                with db() as conn:
                    conn.execute("INSERT INTO files(id,suffix,size,owner_id) VALUES (?,?,?,?)",
                                 (cache_id, suffix, len(downloaded.content), request.state.user_id))
            except BaseException:
                cached.unlink(missing_ok=True)
                raise
            payload["file_id"] = cache_id
        elif reference:
            path = file_path(reference, request.state.user_id)
            if path.suffix not in {".docx", ".pdf"}:
                raise error(422, "TEXT_REQUIRED", "doc 请先转为 docx/pdf，zip 请用于代码测试；也可提交提取后的 content")
            payload["file_id"] = reference
        return enqueue("review", payload, request.state.request_id, request.state.user_id)

    @app.post("/api/code-tests", status_code=202)
    async def code_test(query: CodeQuery, request: Request):
        if bool(query.code) == bool(query.file_id):
            raise error(422, "INVALID_INPUT", "code 和 file_id 必须且只能提供一个")
        payload = query.model_dump(exclude_none=True, exclude={"file_id"})
        if sandbox_protocol == "text" and (query.file_id or not (query.tests.strip() and (query.code or "").strip())):
            raise error(422, "CODE_AND_TESTS_REQUIRED", "当前工程接口仅支持单份Java代码及JUnit测试源码，暂不支持ZIP项目")
        if query.file_id:
            path = file_path(query.file_id, request.state.user_id)
            if path.suffix != ".zip":
                raise error(422, "INVALID_INPUT", "代码项目必须是 zip")
            # Shared volume; the sandbox owns safe extraction and container execution.
            payload["archive_path"] = str(path.resolve())
        return enqueue("code", payload, request.state.request_id, request.state.user_id)

    @app.get("/api/tasks/{task_id}")
    async def task_status(task_id: str, request: Request):
        with db() as conn:
            row = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        if not row or (backend_url and row["owner_id"] != request.state.user_id):
            raise error(404, "TASK_NOT_FOUND", "任务不存在")
        return {"task_id": task_id, "status": row["status"], "progress": row["progress"],
                "request_id": row["request_id"], "result": json.loads(row["result"]) if row["result"] else None,
                "error": json.loads(row["error"]) if row["error"] else None}

    return app


app = create_app()
